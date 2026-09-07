# -*- coding: utf-8 -*-
"""Конвертер docs/Презентация и блоки данных.md -> docs/Презентация и блоки данных_чтение.docx.

Поддержка: заголовки #..####, таблицы |...|, списки "- ", цитаты "> ",
жирный **...**, код `...`, ссылки [текст](url) (выводятся текстом).
Блок ```mermaid ... ``` заменяется примечанием.
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

SRC = "docs/Презентация и блоки данных.md"
OUT = "docs/Презентация и блоки данных_чтение.docx"

LINK_RE = re.compile(r"\[([^\[\]]*)\]\([^)]*\)")
BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")


def add_runs(par, text):
    """Разбирает инлайн-разметку: жирный, код, ссылки (текстом)."""
    depth = 0
    while text and depth < 10:
        m = LINK_RE.search(text)
        if m:
            if m.start():
                add_inline(par, text[: m.start()])
            add_runs(par, m.group(1))
            text = text[m.end():]
            depth += 1
            continue
        add_inline(par, text)
        return


def add_inline(par, text):
    while text:
        m = BOLD_RE.search(text)
        c = CODE_RE.search(text)
        first = None
        if m and (c is None or m.start() <= c.start()):
            first = ("bold", m)
        elif c:
            first = ("code", c)
        if first is None:
            run = par.add_run(text)
            return
        kind, mt = first
        if mt.start():
            par.add_run(text[: mt.start()])
        run = par.add_run(mt.group(1))
        if kind == "bold":
            run.bold = True
        else:
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        text = text[mt.end():]


def normalize_links(s):
    """Ссылки [текст](url) -> текст (упрощение перед разбором)."""
    return LINK_RE.sub(lambda m: m.group(1), s)


def main():
    with open(SRC, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()

    # Базовые настройки страницы и шрифта
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()

        # Блок кода (в документе — только mermaid-схема презентации)
        if line.startswith("```"):
            is_mermaid = line.startswith("```mermaid")
            i += 1
            while i < n and not lines[i].startswith("```"):
                i += 1
            if i < n:
                i += 1  # закрывающий ```
            if is_mermaid:
                p = doc.add_paragraph()
                r = p.add_run("Схема целевой раскладки презентации приведена в таблице ниже.")
                r.italic = True
            continue

        # Горизонтальные линии и пустые строки
        if set(line.strip()) <= set("-") and len(line.strip()) >= 3:
            i += 1
            continue
        if not line.strip():
            i += 1
            continue

        # Заголовки
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            par = doc.add_heading(level=level)
            add_runs(par, normalize_links(m.group(2)))
            i += 1
            continue

        # Таблицы: собираем подряд идущие строки, начинающиеся с '|'
        if line.strip().startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells_rows = []
            for row in rows:
                if set(row.replace("|", "").replace("-", "").replace(":", "").replace(" ", "")) == set():
                    continue  # разделитель |---|---|
                cells = [c.strip() for c in row.strip("|").split("|")]
                cells_rows.append(cells)
            if cells_rows:
                ncols = max(len(r) for r in cells_rows)
                table = doc.add_table(rows=len(cells_rows), cols=ncols)
                table.style = "Table Grid"
                for ri, crow in enumerate(cells_rows):
                    for ci in range(ncols):
                        cell_text = crow[ci] if ci < len(crow) else ""
                        cell = table.cell(ri, ci)
                        cell.text = ""
                        par = cell.paragraphs[0]
                        add_runs(par, normalize_links(cell_text))
            continue

        # Цитаты
        if line.startswith(">"):
            parts = []
            while i < n and lines[i].strip().startswith(">"):
                parts.append(lines[i].strip()[1:].strip())
                i += 1
            par = doc.add_paragraph(style="Quote")
            add_runs(par, normalize_links(" ".join(parts)))
            continue

        # Маркированные списки
        if re.match(r"^-\s+", line):
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, normalize_links(re.sub(r"^-\s+", "", line)))
            i += 1
            continue

        # Обычный абзац
        par = doc.add_paragraph()
        add_runs(par, normalize_links(line))
        i += 1

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = main()
    # Проверка результата: переоткрываем и считаем структуру
    d = Document(path)
    print("saved:", path)
    print("paragraphs:", len(d.paragraphs))
    print("tables:", len(d.tables))
    for idx, t in enumerate(d.tables, 1):
        print("table", idx, "rows", len(t.rows), "cols", len(t.columns))
    sys.exit(0)
